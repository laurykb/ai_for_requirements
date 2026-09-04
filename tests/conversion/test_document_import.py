from api.lynx_corpus import merge_requirements, parse_corpus_payloads
from conversion.document_import import extract_marked_text


SAMPLE = """
3.2.1 Caractéristiques fonctionnelles
[PIDS-CGLSARC-REQ-138]
Le CGSARC doit être fourni avec un kit installation et désinstallation.

[PIDS-CGLSARC-REQ-139]
Le CGSARC doit journaliser les opérations.
"""


def test_extrait_exigences_marquees_avec_provenance_et_section():
    requirements, warnings = extract_marked_text(SAMPLE, "PIDS-CGLSARC.doc")

    assert warnings == []
    assert [req["id"] for req in requirements] == [
        "PIDS-CGLSARC-REQ-138", "PIDS-CGLSARC-REQ-139",
    ]
    first = requirements[0]
    assert first["texte"] == "Le CGSARC doit être fourni avec un kit installation et désinstallation."
    assert first["niveau"] == 1
    assert first["domaine"] == "CGLSARC"
    assert first["occurrences"][0]["section"] == "3.2.1"
    assert first["occurrences"][0]["document_type"] == "PIDS"


def test_doublon_identique_fusionne_les_occurrences_sans_collision():
    requirements, warnings = extract_marked_text(
        SAMPLE + "\n[PIDS-CGLSARC-REQ-138]\nLe CGSARC doit être fourni avec un kit installation et désinstallation.\n",
        "PIDS-CGLSARC.doc",
    )

    first = next(req for req in requirements if req["id"] == "PIDS-CGLSARC-REQ-138")
    assert len(first["occurrences"]) == 2
    assert first["collision_variants"] == []
    assert not any("REQ-138" in warning for warning in warnings)


def test_collision_inter_sources_conserve_les_deux_formulations():
    left, _ = extract_marked_text(SAMPLE, "PIDS-CGLSARC.doc")
    right, _ = extract_marked_text(
        "[PIDS-CGLSARC-REQ-138]\nLe CGSARC doit inclure un kit complet installation.\n",
        "DJEM.xls.txt",
    )

    merged, warnings = merge_requirements([left, right])
    requirement = next(req for req in merged if req["id"] == "PIDS-CGLSARC-REQ-138")

    assert len(requirement["occurrences"]) == 2
    assert len(requirement["collision_variants"]) == 2
    assert any("collision" in warning for warning in warnings)


def test_parseur_accepte_un_fichier_texte_et_preserve_les_metadonnees():
    requirements, warnings = parse_corpus_payloads([
        ("PIDS-CGLSARC.txt", SAMPLE.encode()),
    ])

    assert len(requirements) == 2
    assert requirements[0]["occurrences"][0]["source"] == "PIDS-CGLSARC.txt"
    assert warnings == []


def test_une_nouvelle_section_termine_le_bloc_exigence():
    text = """5 Livraison
[PIDS-CGLSARC-REQ-138]
Le produit doit inclure un kit complet.

6 Annexe
Cette annexe ne fait pas partie de lexigence.
"""

    requirements, _ = extract_marked_text(text, "PIDS-CGLSARC.doc")

    assert requirements[0]["texte"] == "Le produit doit inclure un kit complet."


def test_marqueur_irregulier_et_endreq_produisent_un_texte_canonique():
    requirements, warnings = extract_marked_text(
        "[PIDS-OPL_TMSC- REQ-2]\nLe système doit démarrer.\nendReq\nNota hors exigence.\n",
        "PIDS-OPL-TMSC.doc",
    )

    assert warnings == []
    assert requirements[0]["id"] == "PIDS-OPL_TMSC-REQ-2"
    assert requirements[0]["texte"] == "Le système doit démarrer."


def test_table_word_separe_texte_liens_voisins_et_metadonnees():
    from io import BytesIO
    from docx import Document
    from conversion.document_import import import_document

    document = Document()
    document.add_paragraph("[PIDS-CGT-REQ-001]")
    document.add_paragraph("Le CGT doit journaliser les actions.")
    document.add_paragraph("endReq")
    document.add_paragraph("[PIDS-CGT-REQ-002]")
    document.add_paragraph("Le CGT doit protéger les accès.")
    document.add_paragraph("endReq")
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Exigence"
    table.cell(0, 1).text = "Exigence amont"
    table.cell(1, 0).text = "PIDS-CGT-REQ-001"
    table.cell(1, 1).text = "[SSS-STC-E-REQ-0020]\n[SSS-STC-E-REQ-1097]"
    table.cell(2, 0).text = "PIDS-CGT-REQ-002"
    table.cell(2, 1).text = ""
    stream = BytesIO()
    document.save(stream)

    requirements, warnings = import_document(stream.getvalue(), "PIDS-CGT.docx")
    by_id = {requirement["id"]: requirement for requirement in requirements}

    assert by_id["PIDS-CGT-REQ-001"]["texte"] == "Le CGT doit journaliser les actions."
    assert by_id["PIDS-CGT-REQ-001"]["links"] == [
        {"type": "SATISFIES", "target": "SSS-STC-E-REQ-0020"},
        {"type": "SATISFIES", "target": "SSS-STC-E-REQ-1097"},
    ]
    table_occurrence = next(
        occurrence for occurrence in by_id["PIDS-CGT-REQ-001"]["occurrences"]
        if occurrence.get("kind") == "traceability_table"
    )
    assert table_occurrence["neighbor_ids"] == [
        "SSS-STC-E-REQ-0020", "SSS-STC-E-REQ-1097",
    ]
    assert table_occurrence["table_index"] == 1
    assert table_occurrence["row_index"] == 2
    assert by_id["PIDS-CGT-REQ-001"]["collision_variants"] == []
    assert any("1 matrice(s)" in warning for warning in warnings)


def test_variantes_internes_ne_deviennent_pas_des_divergences_inter_sources():
    requirements, warnings = extract_marked_text(
        "[ABC-REQ-001]\nTexte court mais valide.\n"
        "[ABC-REQ-001]\nTexte canonique sensiblement plus complet et valide.\n",
        "source.doc",
    )
    assert requirements[0]["texte"] == "Texte canonique sensiblement plus complet et valide."
    assert requirements[0]["collision_variants"] == []
    assert any("occurrences internes" in warning for warning in warnings)


def test_bloc_word_style_exigence_conserve_titre_corps_et_liens():
    from io import BytesIO

    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    from conversion.document_import import import_document

    document = Document()
    styles = document.styles
    styles.add_style("Exig_Num", WD_STYLE_TYPE.PARAGRAPH)
    styles.add_style("Orig_Rqt", WD_STYLE_TYPE.PARAGRAPH)
    styles.add_style("Requirement", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("ICD-XSMTP-75: Dépôt d'un message", style="Exig_Num")
    document.add_paragraph("#XSMTP-SMTP-03", style="Orig_Rqt")
    document.add_paragraph("{", style="Requirement")
    document.add_paragraph("Le service doit accepter le message.", style="Requirement")
    document.add_paragraph("}", style="Requirement")
    stream = BytesIO()
    document.save(stream)

    requirements, warnings = import_document(stream.getvalue(), "ICD-XSMTP.docx")

    assert warnings == []
    assert requirements[0]["id"] == "ICD-XSMTP-75"
    assert requirements[0]["texte"] == "Le service doit accepter le message."
    assert requirements[0]["links"] == [
        {"type": "SATISFIES", "target": "XSMTP-SMTP-03"},
    ]
    occurrence = requirements[0]["occurrences"][0]
    assert occurrence["kind"] == "styled_requirement"
    assert occurrence["title"] == "Dépôt d'un message"
    assert occurrence["neighbor_ids"] == ["XSMTP-SMTP-03"]




def test_fusion_meme_source_ne_cree_pas_de_collision():
    base = {
        "id": "ABC-REQ-001", "niveau": 1, "type": "Exigence",
        "domaine": "ABC", "parent_id": None, "test_status": "PENDING",
        "links": [], "source": "source.doc", "occurrences": [],
        "collision_variants": [],
    }
    merged, warnings = merge_requirements([
        [{**base, "texte": "Formulation courte."}],
        [{**base, "texte": "Formulation canonique plus complète."}],
    ])

    assert merged[0]["texte"] == "Formulation canonique plus complète."
    assert merged[0]["collision_variants"] == []
    assert any("occurrences différentes" in warning for warning in warnings)

