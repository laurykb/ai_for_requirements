# tests/test_multiformat_ingest.py
"""Conversion offline .docx et .xlsx -> Markdown non vide via le pipeline Docling.
Fichiers minimaux générés à la volée (python-docx / openpyxl, déjà installés)."""
from pathlib import Path
import pytest
from preprocessing.pdf_to_markdown import convert_and_clean


def _make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Menaces identifiees", level=1)
    d.add_paragraph("Le systeme doit resister a l'attaque par rejeu et a l'injection.")
    d.save(path)


def _make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Categorie", "Attaque"])
    ws.append(["Reseau", "Deni de service"])
    ws.append(["Applicatif", "Injection SQL"])
    wb.save(path)


@pytest.mark.parametrize("maker,ext", [(_make_docx, "docx"), (_make_xlsx, "xlsx")])
def test_office_formats_convert_to_markdown(tmp_path, maker, ext):
    src = tmp_path / f"echantillon.{ext}"
    maker(str(src))
    md_path = convert_and_clean(str(src), out_dir=str(tmp_path / "out"))
    text = Path(md_path).read_text(encoding="utf-8")
    assert text.strip(), f"markdown vide pour .{ext}"
    # Un token du contenu source doit survivre à la conversion.
    needle = "rejeu" if ext == "docx" else "Injection"
    assert needle.lower() in text.lower()
