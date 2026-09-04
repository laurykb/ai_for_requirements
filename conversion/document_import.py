"""Extraction déterministe des exigences marquées dans les documents bureautiques."""
from __future__ import annotations

import re
import io
import shutil
import subprocess
import tempfile
from collections import defaultdict
from pathlib import Path

from docx import Document


MARKER_RE = re.compile(
    r"(?m)^[ \t]*\[([A-Z0-9][A-Z0-9_.-]{1,100}REQ[A-Z0-9_.-]*[0-9]+)\][ \t]*$"
)
LOOSE_MARKER_RE = re.compile(r"^\s*\[([^\]]*REQ[^\]]*\d+\.?)\]\s*$", re.IGNORECASE)
TOKEN_RE = re.compile(
    r"\b[A-Z][A-Z0-9]{1,15}(?:[-_.][A-Z0-9]{1,16}){1,8}[-_.]\d{1,7}\b"
)
HEADING_RE = re.compile(r"^\s*((?:\d+\.)*\d+)\s+(.{2,180})\s*$")
STYLE_MARKER_RE = re.compile(
    r"^\s*([A-Z][A-Z0-9_. -]{1,80}\d)\s*(?::|\t)\s*(.*)$"
)


def _document_kind(filename: str) -> str:
    stem = Path(filename).stem.upper().replace("_", "-")
    for kind in ("DJEM", "SSDD", "SSS", "PIDS", "IRS", "ICD", "NST", "DD"):
        if stem.startswith(kind):
            return kind
    return "DOCUMENT"


def _document_level(kind: str) -> int:
    return {"SSS": 0, "DJEM": 0, "SSDD": 1, "PIDS": 1, "IRS": 1}.get(kind, 2)


def _domain_from_id(ident: str) -> str:
    parts = ident.split("-")
    try:
        req_index = parts.index("REQ")
    except ValueError:
        req_index = next((i for i, part in enumerate(parts) if part.startswith("REQ")), len(parts) - 1)
    component = parts[1:req_index]
    return " / ".join(component) if component else "Général"

def _canonical_id(value: str) -> str:
    """Normalise les espaces parasites et la ponctuation d'un identifiant."""
    return re.sub(r"\s+", "", value).rstrip(".").upper()


def _normalise_marker_lines(text: str) -> str:
    lines = []
    for line in text.replace("\r\n", "\n").replace("\r", "\n").splitlines():
        match = LOOSE_MARKER_RE.fullmatch(line)
        lines.append(f"[{_canonical_id(match.group(1))}]" if match else line)
    return "\n".join(lines)


def _ids(value: str) -> list[str]:
    return list(dict.fromkeys(_canonical_id(match.group(0)) for match in TOKEN_RE.finditer(value)))



def _convert_to_text(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        return data.decode("utf-8-sig", errors="replace")
    if suffix not in {".doc", ".docx", ".odt"}:
        raise ValueError("Format documentaire attendu : .doc, .docx, .odt ou .txt")
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise ValueError("LibreOffice est requis pour lire les documents Word")
    with tempfile.TemporaryDirectory(prefix="lynx-document-") as directory:
        root = Path(directory)
        source = root / f"upload{suffix}"
        source.write_bytes(data)
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "txt:Text", "--outdir", directory, str(source)],
            capture_output=True, text=True, timeout=180, check=False,
        )
        output = root / "upload.txt"
        if result.returncode or not output.exists():
            detail = (result.stderr or result.stdout).strip()
            suffix_detail = f": {detail}" if detail else ""
            raise ValueError(f"Conversion du document impossible{suffix_detail}")
        return output.read_text(encoding="utf-8", errors="replace")


def _section_before(text: str, offset: int) -> tuple[str | None, str | None]:
    number = title = None
    for line in reversed(text[:offset].splitlines()[-80:]):
        match = HEADING_RE.match(line)
        if match:
            number, title = match.group(1), match.group(2).strip()
            break
    return number, title


def _clean_body(value: str, current_section: str | None = None) -> str:
    lines = [line.rstrip() for line in value.replace("\r\n", "\n").replace("\r", "\n").splitlines()]
    end = next((index for index, line in enumerate(lines)
                if line.strip().casefold() == "endreq"), None)
    if end is not None:
        lines = lines[:end]
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()
    for index, line in enumerate(lines[1:], 1):
        heading = HEADING_RE.match(line)
        if heading and heading.group(1) != current_section:
            lines = lines[:index]
            while lines and not lines[-1].strip():
                lines.pop()
            break
    return "\n".join(lines).strip()


def _candidate_score(body: str) -> tuple[int, int]:
    useful = len(re.sub(r"\s+", " ", body).strip())
    plausible = 1 if 15 <= useful <= 8000 else 0
    return plausible, min(useful, 8000)


def _strip_requirement_braces(lines: list[str]) -> str:
    """Retire uniquement les accolades qui délimitent un bloc DOORS exporté."""
    value = "\n".join(lines).strip()
    if value.startswith("{"):
        value = value[1:].lstrip()
    if value.endswith("}"):
        value = value[:-1].rstrip()
    return value


def _extract_styled_requirements(document, filename: str) -> tuple[list[dict], list[str]]:
    """Extrait les blocs Word formels ``Exig_SMTP`` et ``Exig_Num``."""
    paragraphs = list(document.paragraphs)
    marker_indexes: list[tuple[int, re.Match[str]]] = []
    accepted_styles = {"exig_smtp", "exig_num"}
    for index, paragraph in enumerate(paragraphs):
        style = paragraph.style.name if paragraph.style is not None else ""
        if style.casefold() not in accepted_styles:
            continue
        match = STYLE_MARKER_RE.match(paragraph.text.strip())
        if match:
            marker_indexes.append((index, match))

    kind = _document_kind(filename)
    requirements: list[dict] = []
    warnings: list[str] = []
    for position, (index, marker) in enumerate(marker_indexes):
        stop = marker_indexes[position + 1][0] if position + 1 < len(marker_indexes) else len(paragraphs)
        ident = _canonical_id(marker.group(1))
        title = marker.group(2).strip() or None
        body_lines: list[str] = []
        upstream_ids: list[str] = []
        body_started = False
        for paragraph in paragraphs[index + 1:stop]:
            style = paragraph.style.name if paragraph.style is not None else ""
            value = paragraph.text.strip()
            if not value:
                continue
            if style.casefold() == "orig_rqt" and not body_started:
                upstream_ids.extend(_ids(value.replace("#", " ")))
                continue
            if value.startswith("{"):
                body_started = True
            if body_started:
                body_lines.append(value)
                if value.endswith("}"):
                    break
        body = _strip_requirement_braces(body_lines)
        if not body:
            warnings.append(f"{ident}: bloc d'exigence stylé sans énoncé exploitable dans {filename}")
            continue
        occurrence = {
            "source": filename,
            "document_type": kind,
            "kind": "styled_requirement",
            "paragraph_index": index,
            "style": paragraphs[index].style.name,
            "title": title,
            "neighbor_ids": list(dict.fromkeys(upstream_ids)),
        }
        requirements.append({
            "id": ident,
            "niveau": _document_level(kind),
            "type": "Exigence",
            "domaine": _domain_from_id(ident),
            "texte": body,
            "parent_id": None,
            "test_status": "PENDING",
            "links": [
                {"type": "SATISFIES", "target": target}
                for target in occurrence["neighbor_ids"] if target != ident
            ],
            "source": filename,
            "occurrences": [occurrence],
            "collision_variants": [],
        })
    return requirements, warnings
def extract_marked_text(text: str, filename: str) -> tuple[list[dict], list[str]]:
    """Extrait et déduplique les blocs marqués dans un texte normalisé."""
    text = _normalise_marker_lines(text)
    matches = list(MARKER_RE.finditer(text))
    kind = _document_kind(filename)
    candidates: dict[str, list[dict]] = defaultdict(list)
    warnings: list[str] = []
    for index, marker in enumerate(matches):
        ident = marker.group(1)
        stop = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        section, section_title = _section_before(text, marker.start())
        body = _clean_body(text[marker.end():stop], section)
        occurrence = {
            "source": filename,
            "document_type": kind,
            "section": section,
            "section_title": section_title,
            "offset": marker.start(),
        }
        candidates[ident].append({"texte": body, "occurrence": occurrence})

    requirements = []
    for ident, variants in candidates.items():
        non_empty = [item for item in variants if item["texte"]]
        if not non_empty:
            warnings.append(f"{ident}: marqueur sans énoncé exploitable dans {filename}")
            continue
        chosen = max(non_empty, key=lambda item: _candidate_score(item["texte"]))
        distinct = []
        signatures = set()
        for item in non_empty:
            signature = re.sub(r"\s+", " ", item["texte"]).strip().casefold()
            if signature in signatures:
                continue
            signatures.add(signature)
            distinct.append({**item["occurrence"], "texte": item["texte"]})
        if len(distinct) > 1:
            warnings.append(
                f"{ident}: {len(distinct)} occurrences internes différentes dans {filename}; "
                "la formulation canonique la plus complète est retenue"
            )
        requirements.append({
            "id": ident,
            "niveau": _document_level(kind),
            "type": "Exigence",
            "domaine": _domain_from_id(ident),
            "texte": chosen["texte"],
            "parent_id": None,
            "test_status": "PENDING",
            "links": [],
            "source": filename,
            "occurrences": [item["occurrence"] for item in variants],
            "collision_variants": [],
        })
    if not requirements:
        warnings.append(f"{filename}: aucun marqueur exigence détecté")
    return requirements, warnings


def _docx_bytes(data: bytes, filename: str) -> bytes:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return data
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise ValueError("LibreOffice est requis pour lire les documents Word")
    with tempfile.TemporaryDirectory(prefix="lynx-document-") as directory:
        root = Path(directory)
        source = root / f"upload{suffix}"
        source.write_bytes(data)
        result = subprocess.run(
            [executable, f"-env:UserInstallation={root.joinpath('lo-profile').as_uri()}",
             "--headless", "--convert-to", "docx", "--outdir", directory, str(source)],
            capture_output=True, text=True, timeout=180, check=False,
        )
        output = root / "upload.docx"
        if result.returncode or not output.exists():
            detail = (result.stderr or result.stdout).strip()
            suffix_detail = f": {detail}" if detail else ""
            raise ValueError(f"Conversion du document impossible{suffix_detail}")
        return output.read_bytes()


def _add_table_evidence(document, requirements: list[dict], filename: str) -> list[str]:
    """Ajoute liens, voisins et provenance depuis les matrices Word réelles."""
    by_id = {requirement["id"]: requirement for requirement in requirements}
    warnings: list[str] = []
    matrix_count = matrix_rows = 0
    for table_index, table in enumerate(document.tables, start=1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if len(rows) < 2 or max((len(row) for row in rows), default=0) < 2:
            continue
        n_cols = max(len(row) for row in rows)
        column_ids = [
            [ident for cells in rows[1:] if col < len(cells) for ident in _ids(cells[col])]
            for col in range(n_cols)
        ]
        local_scores = [sum(ident in by_id for ident in identifiers) for identifiers in column_ids]
        if not local_scores or max(local_scores) == 0:
            continue
        local_col = max(range(n_cols), key=lambda col: local_scores[col])
        upstream_cols = [
            col for col, identifiers in enumerate(column_ids)
            if col != local_col and any(ident not in by_id for ident in identifiers)
        ]
        if not upstream_cols:
            continue
        matrix_count += 1
        for row_index, cells in enumerate(rows[1:], start=2):
            if local_col >= len(cells):
                continue
            local_ids = _ids(cells[local_col])
            upstream_ids = list(dict.fromkeys(
                ident for col in upstream_cols if col < len(cells) for ident in _ids(cells[col])
            ))
            if not local_ids:
                continue
            matrix_rows += 1
            for local_id in local_ids:
                requirement = by_id.get(local_id)
                if requirement is None:
                    warnings.append(
                        f"{filename}: table {table_index}, ligne {row_index}: "
                        f"{local_id} absent des définitions du document"
                    )
                    continue
                requirement["occurrences"].append({
                    "source": filename,
                    "kind": "traceability_table",
                    "table_index": table_index,
                    "row_index": row_index,
                    "headers": rows[0],
                    "cells": cells,
                    "neighbor_ids": upstream_ids,
                })
                known = {(link["type"], link["target"]) for link in requirement["links"]}
                for upstream_id in upstream_ids:
                    if upstream_id != local_id and ("SATISFIES", upstream_id) not in known:
                        requirement["links"].append({"type": "SATISFIES", "target": upstream_id})
                        known.add(("SATISFIES", upstream_id))
    if matrix_count:
        warnings.append(
            f"{filename}: {matrix_count} matrice(s) de traçabilité structurée(s), "
            f"{matrix_rows} ligne(s) reliée(s)"
        )
    return warnings


def import_document(data: bytes, filename: str) -> tuple[list[dict], list[str]]:
    """Sépare définitions canoniques et preuves de traçabilité tabulaires."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        return extract_marked_text(data.decode("utf-8-sig", errors="replace"), filename)
    if suffix not in {".doc", ".docx", ".odt"}:
        raise ValueError("Format documentaire attendu : .doc, .docx, .odt ou .txt")
    document = Document(io.BytesIO(_docx_bytes(data, filename)))
    paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    requirements, warnings = extract_marked_text(paragraph_text, filename)
    styled_requirements, styled_warnings = _extract_styled_requirements(document, filename)
    existing_ids = {requirement["id"] for requirement in requirements}
    requirements.extend(
        requirement for requirement in styled_requirements
        if requirement["id"] not in existing_ids
    )
    if styled_requirements:
        warnings = [warning for warning in warnings if "aucun marqueur exigence" not in warning]
    warnings.extend(styled_warnings)
    warnings.extend(_add_table_evidence(document, requirements, filename))
    return requirements, warnings
