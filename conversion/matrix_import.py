"""Lecture transversale de matrices de traçabilité XLS/XLSX.

La structure est découverte dans le contenu : chaque cellule contenant des
identifiants est associée à la cellule de texte située à sa droite. Aucun nom
de feuille, préfixe d'identifiant ou intitulé de colonne n'est imposé.
"""
from __future__ import annotations

import io
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

TOKEN_RE = re.compile(
    r"\b[A-Z][A-Z0-9]{1,15}(?:[-_.][A-Z0-9]{1,16}){1,8}[-_.]\d{1,7}\b"
)


def _ids(value: object) -> list[str]:
    """Identifiants uniques trouvés dans une cellule, ordre conservé."""
    if value is None:
        return []
    return list(dict.fromkeys(m.group(0).strip() for m in TOKEN_RE.finditer(str(value))))


def _texts_by_id(value: object, ids: list[str]) -> dict[str, str]:
    """Découpe une cellule `[ID] texte [ID] texte`; repli sur le texte entier."""
    text = str(value or "").strip()
    if not text:
        return {}
    markers = [(m.group(1), m.start(), m.end()) for m in re.finditer(
        r"\[?\s*(" + TOKEN_RE.pattern + r")\s*\]?", text)]
    found: dict[str, str] = {}
    for index, (ident, _start, end) in enumerate(markers):
        if ident not in ids:
            continue
        stop = markers[index + 1][1] if index + 1 < len(markers) else len(text)
        body = text[end:stop].strip(" \n\t:;-–")
        if body:
            found[ident] = body
    if len(ids) == 1 and ids[0] not in found:
        found[ids[0]] = text
    return found


def rows_to_requirements(rows: list[tuple], source: str = "matrice") -> tuple[list[dict], list[str]]:
    """Convertit les lignes d'une ou plusieurs feuilles en exigences LynX."""
    records: dict[str, dict] = {}
    trace: dict[str, set[str]] = {}
    warnings: list[str] = []
    for row_index, row in enumerate(rows, start=1):
        groups: list[tuple[int, list[str]]] = []
        for col, cell in enumerate(row):
            ids = _ids(cell)
            if not ids:
                continue
            # Une colonne de texte répétant ses ids n'est pas un nouveau groupe.
            previous = groups[-1][1] if groups else []
            if previous and set(ids).issubset(previous):
                continue
            groups.append((col, ids))
            text_cell = row[col + 1] if col + 1 < len(row) else ""
            texts = _texts_by_id(text_cell, ids)
            for ident in ids:
                body = texts.get(ident, "").strip()
                if not body:
                    continue
                existing = records.get(ident)
                if existing is None or len(body) > len(existing["texte"]):
                    records[ident] = {
                        "id": ident, "niveau": len(groups) - 1,
                        "type": "Exigence", "domaine": "Général", "texte": body,
                        "parent_id": None, "test_status": "PENDING", "links": [],
                        "source": source,
                        "occurrences": [{"source": source, "row": row_index, "column": col + 1}],
                    }
        # La direction des feuilles peut varier : on conserve une relation
        # transverse SATISFIES, sans fabriquer une hiérarchie parent/enfant.
        for (_, left), (_, right) in zip(groups, groups[1:]):
            for child in right:
                trace.setdefault(child, set()).update(left)

    # Certaines matrices contiennent la même traçabilité dans les deux sens
    # (feuilles « A vers B » puis « B vers A »). On construit un DAG maximal
    # plutôt que de transformer ces répétitions documentaires en cycles métier.
    parent_graph: dict[str, set[str]] = {ident: set() for ident in records}
    candidates = [(child, parent) for child, targets in trace.items()
                  for parent in targets
                  if child in records and parent in records and child != parent]
    candidates.sort(key=lambda edge: (
        records[edge[0]]["niveau"] - records[edge[1]]["niveau"], edge[0], edge[1]),
        reverse=True)

    def reaches(start: str, target: str) -> bool:
        pending, seen = [start], set()
        while pending:
            current = pending.pop()
            if current == target:
                return True
            if current in seen:
                continue
            seen.add(current)
            pending.extend(parent_graph.get(current, ()))
        return False

    rejected = 0
    rejected_examples: list[str] = []
    for child, parent in candidates:
        if reaches(parent, child):
            rejected += 1
            if len(rejected_examples) < 50:
                rejected_examples.append(f"Relation écartée : {child} → {parent} (cycle réciproque)")
            continue
        parent_graph[child].add(parent)
    for ident, targets in parent_graph.items():
        records[ident]["links"] = [
            {"type": "SATISFIES", "target": target} for target in sorted(targets)]
    if rejected:
        warnings.append(
            f"{rejected} relation(s) réciproque(s) écartée(s) pour conserver une traçabilité acyclique.")
        warnings.extend(rejected_examples)
    if not records:
        warnings.append("Aucune paire identifiant / texte détectée dans la matrice.")
    return list(records.values()), warnings


def _xlsx_bytes(data: bytes, suffix: str) -> bytes:
    if suffix == ".xlsx":
        return data
    if suffix != ".xls":
        raise ValueError("Format attendu : .xls ou .xlsx")
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise ValueError("LibreOffice est requis pour lire les fichiers .xls")
    with tempfile.TemporaryDirectory(prefix="lynx-matrix-") as directory:
        root = Path(directory)
        source = root / "upload.xls"
        source.write_bytes(data)
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "xlsx", "--outdir", directory, str(source)],
            capture_output=True, text=True, timeout=120, check=False,
        )
        output = root / "upload.xlsx"
        if result.returncode or not output.exists():
            raise ValueError("Conversion du fichier .xls impossible")
        return output.read_bytes()


def import_matrix(data: bytes, filename: str) -> tuple[list[dict], list[str]]:
    """Lit toutes les feuilles d'une matrice uploadée."""
    suffix = Path(filename).suffix.lower()
    workbook = load_workbook(io.BytesIO(_xlsx_bytes(data, suffix)), read_only=True, data_only=True)
    rows: list[tuple] = []
    for sheet in workbook.worksheets:
        rows.extend(tuple(row) for row in sheet.iter_rows(values_only=True))
    return rows_to_requirements(rows, source=filename)
